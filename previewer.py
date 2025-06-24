import os
from tkinter import Label, Text, Scrollbar, Frame, BOTH, RIGHT, Y, END
from PIL import Image, ImageTk
try:
    from pdf2image import convert_from_path
except ImportError:
    convert_from_path = None
try:
    import docx
except ImportError:
    docx = None

class Previewer:
    def __init__(self, container):
        self.container = container

        for widget in container.winfo_children():
            widget.destroy()

        self.image_label = Label(container, text="Select a file to preview", foreground='#666',
                                   font=('Segoe UI', 13), background='white', anchor='center')
        self.image_label.pack(expand=True, fill='both')

        self.text_frame = None
        self.text_widget = None
        self.scrollbar = None

        self.preview_image = None 

    def show_image_preview(self, filepath):
        self._clear_widgets()
        self.image_label = Label(self.container, background='white')
        self.image_label.pack(expand=True, fill='both')
        try:
            with Image.open(filepath) as img:
                max_w = max(self.container.winfo_width() - 20, 50)
                max_h = max(self.container.winfo_height() - 20, 50)
                img.thumbnail((max_w, max_h))
                self.preview_image = ImageTk.PhotoImage(img)
                self.image_label.config(image=self.preview_image)
        except Exception:
            self.image_label.config(text="Failed to load image preview")

    def show_file_icon_preview(self, ext):
    
        self._clear_widgets()
        icon_text = ""
        if ext == '.pdf':
            icon_text = "\u270E PDF File\nPreview Not Available"
        elif ext in ['.doc', '.docx']:
            icon_text = "\U0001F4C4 DOC File\nPreview Not Available"
        else:
            icon_text = "No Preview"
        self.image_label = Label(self.container, text=icon_text, image='',
                                  compound=None,
                                  font=('Segoe UI', 24), foreground='#333', background='white', anchor='center')
        self.image_label.pack(expand=True, fill='both')

    def show_pdf_preview(self, filepath):
        self._clear_widgets()
        if convert_from_path is None:
            self.image_label = Label(self.container, text="pdf2image not installed.\nCannot preview PDF.",
                                     foreground='red', font=('Segoe UI', 14), background='white')
            self.image_label.pack(expand=True, fill='both')
            return
        try:
            pages = convert_from_path(filepath, first_page=1, last_page=1, dpi=100)
            if not pages:
                raise Exception("No pages found in PDF")
            img = pages[0]
            max_w = max(self.container.winfo_width() - 20, 50)
            max_h = max(self.container.winfo_height() - 20, 50)
            img.thumbnail((max_w, max_h))
            self.preview_image = ImageTk.PhotoImage(img)
            self.image_label = Label(self.container, image=self.preview_image, background='white')
            self.image_label.pack(expand=True, fill='both')
        except Exception as e:
            self.image_label = Label(self.container, text=f"Failed to load PDF preview:\n{e}",
                                     foreground='red', font=('Segoe UI', 12), background='white')
            self.image_label.pack(expand=True, fill='both')

    def show_docx_preview(self, filepath):
        self._clear_widgets()
        if docx is None:
            self.image_label = Label(self.container, text="python-docx not installed.\nCannot preview DOCX.",
                                     foreground='red', font=('Segoe UI', 14), background='white')
            self.image_label.pack(expand=True, fill='both')
            return
        try:
            doc = docx.Document(filepath)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            text_content = '\n'.join(full_text)
        
            self.text_frame = Frame(self.container)
            self.text_frame.pack(fill=BOTH, expand=True)

            self.scrollbar = Scrollbar(self.text_frame)
            self.scrollbar.pack(side=RIGHT, fill=Y)

            self.text_widget = Text(self.text_frame, wrap='word', yscrollcommand=self.scrollbar.set,
                                    background='white', font=('Segoe UI', 11))
            self.text_widget.pack(expand=True, fill=BOTH)
            self.scrollbar.config(command=self.text_widget.yview)

            self.text_widget.insert(END, text_content)
            self.text_widget.config(state='disabled')
        except Exception as e:
            self.image_label = Label(self.container, text=f"Failed to load DOCX preview:\n{e}",
                                     foreground='red', font=('Segoe UI', 12), background='white')
            self.image_label.pack(expand=True, fill='both')

    def clear_preview(self):
        self._clear_widgets()
        self.image_label = Label(self.container, text="Select a file to preview", foreground='#666',
                                   font=('Segoe UI', 13), background='white', anchor='center')
        self.image_label.pack(expand=True, fill='both')

    def _clear_widgets(self):
        # Remove all widgets from container
        for widget in self.container.winfo_children():
            widget.destroy()
        self.preview_image = None
        self.image_label = None
        self.text_widget = None
        self.scrollbar = None
        self.text_frame = None